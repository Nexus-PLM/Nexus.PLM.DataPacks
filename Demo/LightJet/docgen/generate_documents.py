"""Generates the representative document files the NX-200 dataset attaches.

    python generate_documents.py [output-directory]

Every document record in the dataset names a file. This writes real ones - PDF via
reportlab, Word via python-docx, Excel via openpyxl - so the demo has genuine content in
the Vault rather than placeholders. That matters for three reasons beyond looking right:
full-text search has something to find, the translation service has something to convert,
and a drawing or a first article report opens when somebody in the room asks it to.

The content is representative, not certified. Every file says so on its face: these are
demonstration documents for a fictional aeroplane and nothing in them is design data,
approved data, or a real test result.
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent / "seed"))

import lightjet_documents as DOC                                     # noqa: E402
import lightjet_quality as QA                                        # noqa: E402

from reportlab.lib import colors                                      # noqa: E402
from reportlab.lib.enums import TA_CENTER                             # noqa: E402
from reportlab.lib.pagesizes import A4, landscape                     # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import mm                                    # noqa: E402
from reportlab.platypus import (                                      # noqa: E402
    PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

from docx import Document                                             # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                         # noqa: E402
from docx.shared import Pt, RGBColor                                  # noqa: E402

from openpyxl import Workbook                                         # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

COMPANY = "Light Speed Systems"
DOA = "EASA.21J.0742"
DISCLAIMER = (
    "DEMONSTRATION DOCUMENT. Light Speed Systems and the NX-200 are fictional. This file "
    "is representative sample content for a product lifecycle management demonstration. "
    "It is not design data, not approved data, and not a record of any real test, "
    "inspection or certification activity. Do not use it for any engineering, "
    "airworthiness or procurement purpose."
)

ACCENT = colors.HexColor("#1F3A5F")
LIGHT = colors.HexColor("#E8EDF3")


# ─────────────────────────────────────────────────────────────────────────────
# PDF
# ─────────────────────────────────────────────────────────────────────────────

def _styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=base["Title"], fontSize=16, leading=20,
                                textColor=ACCENT, spaceAfter=4),
        "sub": ParagraphStyle("s", parent=base["Normal"], fontSize=9.5, leading=13,
                              textColor=colors.HexColor("#4A5568"), spaceAfter=10),
        "h": ParagraphStyle("h", parent=base["Heading2"], fontSize=11, leading=14,
                            textColor=ACCENT, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("b", parent=base["Normal"], fontSize=9, leading=12.5,
                               spaceAfter=5),
        "small": ParagraphStyle("sm", parent=base["Normal"], fontSize=7.2, leading=9.5,
                                textColor=colors.HexColor("#6B7280")),
        "stamp": ParagraphStyle("st", parent=base["Normal"], fontSize=7.4, leading=9.5,
                                alignment=TA_CENTER,
                                textColor=colors.HexColor("#8B2828")),
    }


def _kv_table(rows, widths=(52 * mm, 118 * mm)):
    table = Table([[Paragraph(f"<b>{k}</b>", _styles()["body"]),
                    Paragraph(str(v), _styles()["body"])] for k, v in rows],
                  colWidths=widths)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), LIGHT),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D2DD")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return table


def _grid_table(header, rows, widths):
    data = [[Paragraph(f"<b>{h}</b>", _styles()["small"]) for h in header]]
    data += [[Paragraph(str(c), _styles()["small"]) for c in row] for row in rows]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D2DD")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    return table


def _frame(canvas, document):
    """Title block and footer on every page, the way a controlled document carries them."""
    canvas.saveState()
    width, height = document.pagesize
    canvas.setStrokeColor(ACCENT)
    canvas.setLineWidth(0.8)
    canvas.rect(14 * mm, 12 * mm, width - 28 * mm, height - 26 * mm)
    canvas.setFont("Helvetica-Bold", 7.5)
    canvas.setFillColor(ACCENT)
    canvas.drawString(16 * mm, height - 11 * mm,
                      f"{COMPANY}  |  NX-200 Light Jet  |  DOA {DOA}")
    canvas.setFont("Helvetica", 6.6)
    canvas.setFillColor(colors.HexColor("#6B7280"))
    canvas.drawRightString(width - 16 * mm, height - 11 * mm,
                           "UNCONTROLLED WHEN PRINTED")
    canvas.drawString(16 * mm, 8.5 * mm, DISCLAIMER[:118] + "...")
    canvas.drawRightString(width - 16 * mm, 8.5 * mm, f"Page {document.page}")
    canvas.restoreState()


def write_pdf(path: Path, title: str, subtitle: str, blocks, landscape_page=False):
    size = landscape(A4) if landscape_page else A4
    document = SimpleDocTemplate(
        str(path), pagesize=size,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
        title=title, author=COMPANY, subject=subtitle)
    style = _styles()
    story = [Paragraph(title, style["title"]), Paragraph(subtitle, style["sub"])]
    story += blocks
    story += [Spacer(1, 8 * mm), Paragraph(DISCLAIMER, style["stamp"])]
    document.build(story, onFirstPage=_frame, onLaterPages=_frame)


# ─────────────────────────────────────────────────────────────────────────────
# Word
# ─────────────────────────────────────────────────────────────────────────────

def write_docx(path: Path, title: str, subtitle: str, meta_rows, sections):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    header = document.sections[0].header.paragraphs[0]
    header.text = f"{COMPANY}  |  NX-200 Light Jet  |  DOA {DOA}"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "UNCONTROLLED WHEN PRINTED - demonstration document, fictional content"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading = document.add_heading(title, level=0)
    heading.runs[0].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    lead = document.add_paragraph(subtitle)
    lead.runs[0].italic = True

    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, value in meta_rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for section_title, paragraphs in sections:
        document.add_heading(section_title, level=1)
        for text in paragraphs:
            if isinstance(text, tuple):           # a nested table
                head, rows = text
                nested = document.add_table(rows=1, cols=len(head))
                nested.style = "Light List Accent 1"
                for i, column in enumerate(head):
                    nested.rows[0].cells[i].text = str(column)
                for row in rows:
                    cells = nested.add_row().cells
                    for i, value in enumerate(row):
                        cells[i].text = str(value)
            elif text.startswith("- "):
                document.add_paragraph(text[2:], style="List Bullet")
            else:
                document.add_paragraph(text)

    document.add_heading("Statement on this document", level=1)
    note = document.add_paragraph(DISCLAIMER)
    note.runs[0].font.color.rgb = RGBColor(0x8B, 0x28, 0x28)
    note.runs[0].font.size = Pt(8)

    document.save(str(path))


# ─────────────────────────────────────────────────────────────────────────────
# Excel
# ─────────────────────────────────────────────────────────────────────────────

HEAD_FILL = PatternFill("solid", start_color="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=9)
CELL_FONT = Font(size=9)
THIN = Side(style="thin", color="C9D2DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet(workbook, name, header, rows, widths, title=None):
    sheet = workbook.create_sheet(name) if workbook.sheetnames != ["Sheet"] \
        else workbook.active
    sheet.title = name
    row_index = 1
    if title:
        sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=11,
                                                             color="1F3A5F")
        sheet.cell(row=2, column=1,
                   value=f"{COMPANY} - NX-200 Light Jet - DOA {DOA} - "
                         f"DEMONSTRATION DOCUMENT, fictional content").font = Font(
            size=8, italic=True, color="8B2828")
        row_index = 4
    for column, value in enumerate(header, start=1):
        cell = sheet.cell(row=row_index, column=column, value=value)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for offset, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index + offset, column=column, value=value)
            cell.font, cell.border = CELL_FONT, BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = sheet.cell(row=row_index + 1, column=1)
    return sheet


# ─────────────────────────────────────────────────────────────────────────────
# Per-kind builders
# ─────────────────────────────────────────────────────────────────────────────

def build_drawing(path: Path, record: dict):
    """A drawing sheet: title block, revision history and characteristic schedule."""
    attrs = record["attrs"]
    style = _styles()
    blocks = [
        _kv_table([
            ("Drawing number", record["pn"]),
            ("Title", record["description"]),
            ("Issue", record.get("revision", "A")),
            ("Status", record["status"]),
            ("Drawing type", attrs.get("drawingType", "-")),
            ("Sheet size / count", f"{attrs.get('sheetSize', '-')} / "
                                   f"{attrs.get('sheetCount', '-')}"),
            ("Scale / projection", f"{attrs.get('drawingScale', '-')} / "
                                   f"{attrs.get('projectionAngle', '-')}"),
            ("Tolerancing", "ASME Y14.5-2018 applied"
             if attrs.get("gdtApplied") else "Not applied"),
            ("CAD source", attrs.get("cadSourceSystem", "-")),
            ("Drawn / checked / approved",
             f"{attrs.get('authorName', '-')} / {attrs.get('checkedBy', '-')} / "
             f"{attrs.get('approvedBy', '-')}"),
            ("Regulation paragraphs", attrs.get("complianceRegulations", "-")),
            ("Ballooned for FAI",
             f"Yes, issue {attrs.get('balloonedIssue')}"
             if attrs.get("balloonedForFai") else "No"),
            ("Characteristics", attrs.get("characteristicCount", "-")),
            ("Export classification", "Export controlled"
             if attrs.get("confidentiality") == "Export controlled" else "Not controlled"),
        ]),
        Spacer(1, 5 * mm),
        Paragraph("Revision history", style["h"]),
        _grid_table(
            ["Issue", "Date", "Description of change", "Authority", "Approved"],
            ([["A", "2026-04-20", "First issue.", "Original design",
               attrs.get("approvedBy", "-")]]
             if record.get("revision", "A") == "A" else
             [["A", "2026-04-20", "First issue.", "Original design",
               attrs.get("approvedBy", "-")],
              ["B", "2026-09-28",
               "Change implemented. See the engineering change order on the object's "
               "change tab for the reason, the affected items and the effectivity.",
               "Engineering change order", attrs.get("approvedBy", "-")]]),
            [16 * mm, 22 * mm, 74 * mm, 32 * mm, 26 * mm]),
        Spacer(1, 5 * mm),
        Paragraph("Notes", style["h"]),
        Paragraph(
            "1. Interpret this drawing in accordance with ASME Y14.5-2018.<br/>"
            "2. Break all sharp edges 0.2 to 0.4 mm unless otherwise stated.<br/>"
            "3. Material and process specifications are as called out on the item "
            "record; this sheet does not restate them.<br/>"
            "4. Dimensions are in millimetres.<br/>"
            "5. This sheet carries no geometry. It is a representative title-block and "
            "schedule rendering for demonstration purposes.", style["body"]),
    ]
    if attrs.get("balloonedForFai"):
        count = attrs.get("characteristicCount") or 12
        rows = []
        for n in range(1, min(count, 18) + 1):
            rows.append([
                str(n),
                ["Linear dimension", "Diameter", "Position", "Flatness",
                 "Surface finish", "Thread", "Material", "Process"][n % 8],
                ["25.40 +/-0.10", "12.70 H7", "0.20 M A B C", "0.05", "1.6 Ra",
                 "1/4-28 UNJF-3B", "Per item record", "Per process spec"][n % 8],
                "Yes" if n % 5 == 0 else "No",
                "CMM" if n % 3 else "Visual",
            ])
        blocks += [
            PageBreak(),
            Paragraph("Characteristic schedule (ballooned)", style["h"]),
            Paragraph(
                f"Issue {attrs.get('balloonedIssue')} is the ballooned issue. The balloon "
                f"numbers below are the characteristic numbers an AS9102 Form 3 reports "
                f"against, and they must agree. "
                f"{count} characteristics in total; the first "
                f"{min(count, 18)} are listed.", style["body"]),
            _grid_table(
                ["Balloon", "Characteristic", "Requirement", "Key", "Method"],
                rows, [18 * mm, 36 * mm, 48 * mm, 16 * mm, 24 * mm]),
        ]
    write_pdf(path, f"{record['pn']} issue {record.get('revision', 'A')}",
              record["description"], blocks)


def build_generic_docx(path: Path, record: dict, kind: str, sections):
    attrs = record["attrs"]
    meta = [
        ("Document number", record["pn"]),
        ("Title", record["description"]),
        ("Issue", record.get("revision", "A")),
        ("Status", record["status"]),
        ("Document kind", kind),
        ("Author", attrs.get("authorName", "-")),
        ("Checked", attrs.get("checkedBy", "-")),
        ("Approved", attrs.get("approvedBy", "-")),
        ("Issue date", attrs.get("issueDate", "-")),
        ("Retention class", attrs.get("retentionClass", "-")),
        ("Confidentiality", attrs.get("confidentiality", "-")),
    ]
    for label, key in (("Certification basis", "certificationBasis"),
                       ("Regulation paragraphs", "complianceRegulations"),
                       ("Means of compliance", "meansOfCompliance"),
                       ("Approval route", "approvalRoute"),
                       ("DOA approval", "doaApprovalReference"),
                       ("Design assurance level", "designAssuranceLevel"),
                       ("AS9100D clause", "qmsClause"),
                       ("Owning function", "recordOwnerFunction")):
        if attrs.get(key):
            meta.append((label, attrs[key]))
    write_docx(path, record["pn"], record["description"], meta, sections)


def sections_for(record: dict) -> list:
    """Section content shaped by what the document is."""
    attrs = record["attrs"]
    notes = record.get("notes") or ""
    common_purpose = [
        f"This document is {record['description'][0].lower()}"
        f"{record['description'][1:]} for the NX-200 light jet.",
        f"It is issued at revision {record.get('revision', 'A')} and is currently "
        f"{record['status']}.",
    ]
    sections = [("Purpose and scope", common_purpose)]

    if notes:
        sections.append(("Why this document reads as it does", [notes]))

    if attrs.get("assessmentType"):
        sections += [
            ("Method", [
                f"Performed as a {attrs['assessmentType']} in accordance with "
                f"{attrs.get('arpReference', 'ARP4754A')}.",
                f"Function analysed: {attrs.get('functionAnalysed', '-')}.",
            ]),
            ("Result", [
                (["Item", "Value"], [
                    ["Worst case failure effect",
                     attrs.get("worstCaseFailureEffect", "-")],
                    ["Quantitative target", attrs.get("quantitativeTarget", "-")],
                    ["Achieved", attrs.get("achievedProbability", "-")],
                    ["Allocated design assurance level",
                     attrs.get("allocatedDal", "-")],
                    ["Derived requirements raised",
                     attrs.get("derivedRequirementCount", 0)],
                    ["Open hazards", attrs.get("openHazardCount", 0)],
                ]),
                "A derived requirement is one that does not trace up to a higher level "
                "requirement. ARP4754A requires each of them to be fed back into this "
                "assessment, which is why they are counted here and flagged on the "
                "requirement itself rather than left to be rediscovered.",
            ]),
        ]

    if attrs.get("testResult"):
        sections += [
            ("Test configuration", [
                (["Item", "Value"], [
                    ["Standard", attrs.get("do160Revision", "-")],
                    ["Sections", attrs.get("do160Sections", "-")],
                    ["Categories", attrs.get("environmentalCategories", "-")],
                    ["Test plan", attrs.get("testPlanReference", "-")],
                    ["Laboratory", attrs.get("testLaboratory", "-")],
                    ["Accreditation", attrs.get("labAccreditation", "-")],
                    ["Specimens", attrs.get("specimenSerialNumbers", "-")],
                    ["Period", f"{attrs.get('testStartDate', '-')} to "
                               f"{attrs.get('testEndDate', '-')}"],
                ]),
            ]),
            ("Result", [
                f"Outcome: {attrs['testResult']}.",
                f"Deviations taken: {attrs.get('deviationReferences') or 'none'}.",
                f"Retest required: {'yes' if attrs.get('retestRequired') else 'no'}.",
                "A pass with a deviation is a different certification position from a "
                "clean pass. The deviation is named here and its acceptance rationale "
                "is recorded on the compliance item, not buried in a paragraph.",
            ]),
        ]

    if attrs.get("planType"):
        sections += [
            ("Assurance basis", [
                (["Item", "Value"], [
                    ["Plan type", attrs["planType"]],
                    ["Standard", attrs.get("assuranceStandard", "-")],
                    ["Level", attrs.get("applicableLevel", "-")],
                    ["Objectives applicable", attrs.get("objectivesApplicable", "-")],
                    ["Independence required",
                     "yes" if attrs.get("objectivesWithIndependence") else "no"],
                    ["Supplements", attrs.get("supplementsApplied") or "none"],
                    ["Authority agreement", attrs.get("authorityAgreementRef", "-")],
                ]),
                "The objective count and the independence requirement are held as data "
                "on this document so the plan set can be checked for completeness "
                "without reading every plan.",
            ]),
        ]

    if attrs.get("interfaceType"):
        sections += [
            ("Interface definition", [
                (["Item", "Value"], [
                    ["Interface type", attrs["interfaceType"]],
                    ["Party A", attrs.get("partyA", "-")],
                    ["Party B", attrs.get("partyB", "-")],
                    ["Protocol", attrs.get("busProtocol", "-")],
                    ["Baseline", attrs.get("interfaceBaselineRef", "-")],
                    ["Bilaterally agreed",
                     "yes" if attrs.get("bilateralAgreementSigned") else "no"],
                    ["Supplier agreement date",
                     attrs.get("supplierAgreementDate") or "not agreed"],
                ]),
                "Both sides of an interface reference this one controlled definition. "
                "Two drawings that each describe the same interface drift; one document "
                "that both reference cannot.",
            ]),
        ]

    if attrs.get("specificationClass"):
        sections += [
            ("Requirements", [
                f"Specification class: {attrs['specificationClass']}.",
                f"Verification method: {attrs.get('verificationMethod', '-')}.",
                f"Supplier flowdown required: "
                f"{'yes' if attrs.get('flowdownRequired') else 'no'}.",
                "- Requirements are stated as measurable values with tolerances, not as "
                "references to another document.",
                "- Where a numeric criterion comes from the certification basis it is "
                "written out in full, so a supplier cannot read it out of a reference "
                "and get it wrong.",
                "- Every requirement in this specification traces to a requirement "
                "object in the product lifecycle management system.",
            ]),
        ]

    if attrs.get("certificationProgramme"):
        sections += [
            ("Certification basis", [
                (["Item", "Value"], [
                    ["Programme", attrs["certificationProgramme"]],
                    ["Specification", attrs.get("applicableSpecification", "-")],
                    ["Amendment level", attrs.get("amendmentLevel", "-")],
                    ["Authority", attrs.get("certificationAuthority", "-")],
                    ["Planned certification",
                     attrs.get("plannedCertificationDate", "-")],
                ]),
            ]),
            ("Special conditions, findings and deviations", [
                f"Special conditions: {attrs.get('specialConditions') or 'none'}.",
                f"Equivalent safety findings: "
                f"{attrs.get('equivalentSafetyFindings') or 'none'}.",
                f"Deviations requested: {attrs.get('deviationsRequested') or 'none'}.",
                f"Certification review items: "
                f"{attrs.get('certificationReviewItems') or 'none'}.",
            ]),
            ("Design organisation privileges used", [
                attrs.get("doaPrivilegesUsed", "-"),
                f"Level of involvement agreed with the authority: "
                f"{'yes' if attrs.get('levelOfInvolvementAgreed') else 'no'}.",
                "Agreeing the level of involvement is what lets most compliance items "
                "be closed under the design organisation privilege instead of item by "
                "item with the authority. On a Part 21 Subpart J programme it is the "
                "single most schedule-relevant agreement there is.",
            ]),
        ]

    if attrs.get("processFamily"):
        sections += [
            ("Process control", [
                (["Item", "Value"], [
                    ["Process family", attrs["processFamily"]],
                    ["Governing standard", attrs.get("governingStandard", "-")],
                    ["NADCAP accreditation required",
                     "yes" if attrs.get("nadcapAccreditationRequired") else "no"],
                    ["Approved sources", attrs.get("approvedSourceList", "-")],
                    ["Qualified operator required",
                     "yes" if attrs.get("qualifiedOperatorRequired") else "no"],
                    ["Frozen process plan",
                     "yes" if attrs.get("frozenPlanRequired") else "no"],
                    ["Pyrometry class", attrs.get("pyrometryClass") or "not applicable"],
                    ["Record retention",
                     f"{attrs.get('recordRetentionYears', '-')} years"],
                ]),
            ]),
            ("Parameter control limits", [
                attrs.get("parameterControlLimits", "-"),
                f"Periodic verification: {attrs.get('periodicTestRequirement') or 'none'}.",
                "A special process is one whose result cannot be verified by inspecting "
                "the finished part. The process itself is therefore the thing under "
                "control, which is why the limits above are stated as numbers.",
            ]),
        ]

    if attrs.get("bulletinType"):
        sections += [
            ("Planning information", [
                (["Item", "Value"], [
                    ["Bulletin type", attrs["bulletinType"]],
                    ["Category", attrs.get("bulletinCategory", "-")],
                    ["Effectivity", attrs.get("effectivityRange", "-")],
                    ["Compliance period", attrs.get("compliancePeriod", "-")],
                    ["Estimated labour",
                     f"{attrs.get('estimatedManHours', '-')} hours"],
                    ["Aircraft downtime",
                     f"{attrs.get('aircraftDowntimeHours', '-')} hours"],
                    ["Modification kit", attrs.get("kitPartNumber") or "none"],
                    ["Weight change", f"{attrs.get('weightChangeKg', 0)} kg"],
                    ["Related directive",
                     attrs.get("relatedDirectiveRef") or "none"],
                    ["Warranty applicable",
                     "yes" if attrs.get("warrantyApplicable") else "no"],
                ]),
            ]),
            ("Accomplishment instructions", [
                "1. Make the aeroplane safe and gain access to the affected area.",
                "2. Remove and discard the superseded part. Record the part number and "
                "serial or batch number removed.",
                "3. Install the replacement part from the modification kit. Torque in "
                "accordance with the installation specification.",
                "4. Carry out the functional check called out in the maintenance manual "
                "chapter for the affected system.",
                "5. Record accomplishment of this bulletin in the aeroplane records, "
                "including the authorised release certificate number of the part "
                "installed.",
                "6. Update the aeroplane's configuration record. An operator who cannot "
                "say which revision is installed cannot show compliance.",
            ]),
        ]

    if attrs.get("complianceStatus"):
        sections += [
            ("Compliance demonstration", [
                (["Item", "Value"], [
                    ["Regulation paragraph", attrs.get("regulationParagraph", "-")],
                    ["Requirement", attrs.get("requirementText", "-")],
                    ["Consensus standard", attrs.get("consensusStandardRef", "-")],
                    ["Means of compliance", attrs.get("meansOfCompliance", "-")],
                    ["Status", attrs.get("complianceStatus", "-")],
                    ["Evidence", attrs.get("evidenceDocumentRefs", "-")],
                    ["Accepted by", attrs.get("acceptedBy") or "not yet accepted"],
                    ["Acceptance date", attrs.get("acceptanceDate") or "-"],
                    ["Open findings", attrs.get("openFindingCount", 0)],
                    ["Authority involvement", attrs.get("levelOfInvolvement", "-")],
                ]),
            ]),
        ]

    if attrs.get("equipmentUnderTest") and attrs.get("assessmentResult"):
        sections += [
            ("Electromagnetic compatibility assessment", [
                (["Item", "Value"], [
                    ["Equipment", attrs["equipmentUnderTest"]],
                    ["Sections assessed", attrs.get("do160Sections", "-")],
                    ["Category", attrs.get("equipmentCategory", "-")],
                    ["HIRF environment", attrs.get("hirfEnvironment", "-")],
                    ["Lightning zone", attrs.get("lightningZone", "-")],
                    ["Laboratory", attrs.get("testLaboratory", "-")],
                    ["Assessment date", attrs.get("assessmentDate", "-")],
                    ["Result", attrs.get("assessmentResult", "-")],
                ]),
                f"Installation limitations: "
                f"{attrs.get('installationLimitations') or 'none'}.",
            ]),
        ]

    sections.append(("Traceability", [
        "Every statement in this document is traceable to an object in the product "
        "lifecycle management system: the requirement it answers, the item it covers, "
        "the compliance item it feeds and the change that last altered it. The object's "
        "own tabs are the master; this file is a rendering.",
    ]))
    return sections


def build_fai_xlsx(path: Path, record: dict):
    """An AS9102 Rev C first article report: Forms 1, 2 and 3 on three sheets."""
    attrs = record["attrs"]
    workbook = Workbook()

    _sheet(workbook, "Form 1 Part Accountability",
           ["Field", "Entry"],
           [["1. Part number", attrs.get("faiPartNumber", "-")],
            ["2. Part name", record["description"]],
            ["3. Serial / lot number", attrs.get("serialOrLotNumber", "-")],
            ["4. FAI report number", record["pn"]],
            ["5. Part revision level", attrs.get("faiPartRevision", "-")],
            ["6. Drawing number", attrs.get("balloonedDrawingRef", "-")],
            ["7. Drawing revision level",
             (attrs.get("balloonedDrawingRef") or "").split()[-1]],
            ["8. Additional changes", "None"],
            ["9. Manufacturing process reference",
             attrs.get("manufacturingProcessRef", "-")],
            ["10. Organisation name", COMPANY],
            ["11. Supplier code", "LSS-K0001"],
            ["12. Purchase order number", "Internal build"],
            ["13. Detail FAI / Assembly FAI",
             "Assembly FAI" if "assembly" in record["description"].lower()
             else "Detail FAI"],
            ["14. Full FAI / Partial FAI", attrs.get("faiType", "-")],
            ["15. Reason for partial or delta FAI", attrs.get("faiReason", "-")],
            ["16. AS9102 revision", attrs.get("as9102Revision", "Rev C")],
            ["17. Characteristics verified", attrs.get("characteristicCount", "-")],
            ["18. Nonconformances found", attrs.get("nonconformanceCount", 0)],
            ["19. Measurement system analysis",
             attrs.get("measurementSystemAnalysisRef", "-")],
            ["20. FAI result", attrs.get("faiResult", "-")],
            ["21. Customer approval required",
             "Yes" if attrs.get("customerApprovalRequired") else "No"],
            ["22. Customer approval date",
             attrs.get("customerApprovalDate") or "Not applicable"],
            ["23. Prepared by", attrs.get("authorName", "-")],
            ["24. Date", attrs.get("issueDate", "-")],
            ["25. Form 3 complete",
             "Yes" if attrs.get("form3Complete") else
             "NO - submission incomplete"]],
           [42, 62],
           title=f"AS9102 {attrs.get('as9102Revision', 'Rev C')} Form 1 - "
                 f"Part Number Accountability - {record['pn']}")

    _sheet(workbook, "Form 2 Product Accountability",
           ["Item", "Material or process name", "Specification", "Code / source",
            "Certificate of conformance", "Status"],
           [[1, "Raw material", "Per item record", "Approved mill",
             "Material certificate on file", "Accepted"],
            [2, "Heat treatment", "AMS 2770 / per item record",
             "NADCAP accredited source", "Process certificate on file", "Accepted"],
            [3, "Chemical conversion coating", "MIL-DTL-5541 Class 1A",
             "NADCAP accredited source", "Process certificate on file", "Accepted"],
            [4, "Fluorescent penetrant inspection", "ASTM E1417 / NAS 410",
             "NAS 410 Level 2 inspector", "Inspection record on file", "Accepted"],
            [5, "Welding", "AWS D17.1 Class A", "Qualified operator, frozen plan",
             "Weld record on file", "Accepted"],
            [6, "Functional test", "Per test specification", "In-house",
             "Test record on file", "Accepted"],
            ["", "Form 2 complete",
             "Yes" if attrs.get("form2Complete") else "NO", "", "", ""]],
           [8, 34, 32, 26, 30, 14],
           title=f"AS9102 Form 2 - Product Accountability - {record['pn']}")

    count = attrs.get("characteristicCount") or 10
    rows = []
    for n in range(1, min(count, 30) + 1):
        nominal = ["25.40", "12.70", "0.20", "0.05", "1.6", "8.00", "3.18", "45.00"][n % 8]
        rows.append([
            n,
            ["Linear dimension", "Diameter", "Position", "Flatness", "Surface finish",
             "Radius", "Thickness", "Angle"][n % 8],
            f"{nominal} +/-0.10",
            f"{float(nominal) + (0.03 if n % 3 else -0.02):.2f}",
            "Pass",
            "CMM" if n % 3 else "Visual",
            "Yes" if n % 5 == 0 else "No",
            "",
        ])
    if attrs.get("nonconformanceCount"):
        rows[-1] = [len(rows), "Radius", "2.50 min", "2.30", "FAIL", "CMM", "Yes",
                    "Raised as a nonconformance; dispositioned use-as-is on the "
                    "strength of the stress re-check. See the nonconformance record."]
    rows.append(["", "Form 3 complete",
                 "Yes" if attrs.get("form3Complete") else "NO", "", "", "", "", ""])

    _sheet(workbook, "Form 3 Characteristics",
           ["Char. no.", "Characteristic", "Requirement", "Result", "Pass / fail",
            "Method", "Key char.", "Nonconformance"],
           rows, [10, 24, 20, 12, 12, 12, 11, 44],
           title=f"AS9102 Form 3 - Characteristic Accountability - {record['pn']} - "
                 f"balloon numbers agree with {attrs.get('balloonedDrawingRef', '-')}")

    _sheet(workbook, "About", ["Note"],
           [[DISCLAIMER],
            ["Forms 1, 2 and 3 are tracked separately on the record in the product "
             "lifecycle management system, because a first article submission is not "
             "complete until all three are, and a delta FAI legitimately re-uses forms "
             "from the full FAI it builds on."],
            ["The balloon numbers on Form 3 must agree with the ballooned issue of the "
             "drawing named on Form 1. 'The latest drawing' is not an answer an auditor "
             "accepts, which is why the drawing issue is recorded explicitly."]],
           [150])

    workbook.save(str(path))


def build_checklist_xlsx(path: Path, record: dict):
    """The compliance checklist: one row per compliance item, from the dataset itself."""
    workbook = Workbook()
    rows = []
    for item in DOC.COMPLIANCE_ITEMS:
        attrs = item["attrs"]
        rows.append([
            attrs.get("regulationParagraph", "-"),
            attrs.get("requirementText", "-"),
            attrs.get("consensusStandardRef", "-"),
            attrs.get("meansOfCompliance", "-"),
            item["pn"],
            item["status"],
            attrs.get("complianceStatus", "-"),
            attrs.get("evidenceDocumentRefs", "-"),
            attrs.get("acceptedBy") or "",
            attrs.get("acceptanceDate") or "",
            attrs.get("openFindingCount", 0),
            attrs.get("levelOfInvolvement", "-"),
        ])
    _sheet(workbook, "Compliance Checklist",
           ["Paragraph", "Requirement", "Consensus standard", "Means of compliance",
            "Compliance item", "Lane", "Compliance status", "Evidence", "Accepted by",
            "Accepted", "Open findings", "Authority involvement"],
           rows, [14, 48, 22, 26, 20, 16, 22, 40, 13, 12, 11, 30],
           title=f"NX-200 compliance checklist - CS-23 Amendment 5 - {record['pn']} "
                 f"revision {record.get('revision', 'A')}")
    _sheet(workbook, "About", ["Note"],
           [[DISCLAIMER],
            ["Every row of this checklist is a compliance item object in the product "
             "lifecycle management system. This spreadsheet is a rendering of that data, "
             "not the master. Change the object and the next rendering changes; change "
             "the spreadsheet and nothing happens, which is the point."],
            [f"Generated {date.today().isoformat()} from "
             f"{len(DOC.COMPLIANCE_ITEMS)} compliance items."]],
           [150])
    workbook.save(str(path))


def build_supplier_xlsx(path: Path, record: dict):
    attrs = record["attrs"]
    workbook = Workbook()
    _sheet(workbook, "Supplier Approval", ["Field", "Entry"],
           [["Supplier", attrs.get("supplierName", "-")],
            ["CAGE / NSCM code", attrs.get("supplierCageCode", "-")],
            ["Approval scope", attrs.get("approvalScope", "-")],
            ["AS9100 certificate", attrs.get("as9100CertificateNumber", "-")],
            ["AS9100 certificate expiry", attrs.get("as9100CertificateExpiry", "-")],
            ["NADCAP accredited scopes", attrs.get("nadcapScopes", "-")],
            ["Last audit", attrs.get("auditDate", "-")],
            ["Audit type", attrs.get("auditType", "-")],
            ["Audit result", attrs.get("auditResult", "-")],
            ["Open findings", attrs.get("openFindingCount", 0)],
            ["Next audit due", attrs.get("nextAuditDue", "-")],
            ["Risk rating", attrs.get("riskRating", "-")],
            ["Approval status", attrs.get("approvalStatus", "-")],
            ["Requirement flowdown accepted",
             "Yes" if attrs.get("flowdownAccepted") else "No"],
            ["Delivered quality", f"{attrs.get('deliveredQualityPpm', '-')} ppm"],
            ["On time delivery", f"{attrs.get('onTimeDeliveryPercent', '-')} %"]],
           [34, 70],
           title=f"Supplier qualification record - {attrs.get('supplierName', '-')} - "
                 f"{record['pn']}")
    _sheet(workbook, "About", ["Note"],
           [[DISCLAIMER],
            ["An expiring AS9100 certificate is a data question, not a diary reminder. "
             "Holding the expiry on the record is what lets a report find every supplier "
             "whose certificate lapses inside ninety days."]],
           [150])
    workbook.save(str(path))


def build_certificate_pdf(path: Path, record: dict, kind: str):
    """A certificate of conformity or an authorised release certificate."""
    attrs = record["attrs"]
    style = _styles()
    if kind == "coc":
        rows = [
            ("Supplier", attrs.get("supplierName", "-")),
            ("CAGE / NSCM code", attrs.get("supplierCageCode", "-")),
            ("Purchase order", f"{attrs.get('purchaseOrderNumber', '-')} line "
                               f"{attrs.get('purchaseOrderLine', '-')}"),
            ("Part number", attrs.get("certifiedPartNumber", "-")),
            ("Part revision", attrs.get("certifiedPartRevision", "-")),
            ("Quantity shipped", attrs.get("quantityShipped", "-")),
            ("Lot / batch", attrs.get("lotNumber", "-")),
            ("Serial numbers", attrs.get("serialNumbers", "-")),
            ("Material certificates", attrs.get("materialCertificateRefs", "-")),
            ("Special process certificates",
             attrs.get("specialProcessCertRefs", "-")),
            ("Deviations declared", attrs.get("deviationsDeclared", "-")),
            ("Authorised signatory", attrs.get("authorisedSignatory", "-")),
            ("Date of issue", attrs.get("issueDate", "-")),
            ("Status", record["status"]),
        ]
        statement = attrs.get("statementOfConformity", "-")
        heading = "Certificate of Conformity"
    else:
        rows = [
            ("Form type", attrs.get("formType", "-")),
            ("3. Form tracking number", attrs.get("certificateNumber", "-")),
            ("4. Organisation name and address",
             f"{COMPANY}, Approval {attrs.get('releasingOrganisationApproval', '-')}"),
            ("6. Item", "1"),
            ("7. Description", record["description"]),
            ("8. Part number", attrs.get("releasedPartNumber", "-")),
            ("Part revision", attrs.get("releasedPartRevision", "-")),
            ("9. Quantity", attrs.get("quantity", "-")),
            ("10. Serial / batch number", attrs.get("serialOrBatchNumber", "-")),
            ("11. Status / work", attrs.get("releaseStatus", "-")),
            ("12. Remarks", attrs.get("remarks", "-")),
            ("Approved design data", attrs.get("approvalDesignData", "-")),
            ("Eligibility", attrs.get("eligibilityStatement", "-")),
            ("Non-approved design data",
             "YES - see block 12" if attrs.get("nonApprovedDesignData") else "No"),
            ("14a. Authorised release",
             attrs.get("authorisedSignatory", "-")),
            ("14e. Date", attrs.get("dateOfIssue", "-")),
        ]
        statement = (
            "Certifies that unless otherwise specified in block 12, the work identified "
            "and described above was accomplished in accordance with the applicable "
            "requirements and in respect of that work the item is considered ready for "
            "release to service."
            + (" NOTE: this release is against non-approved design data, so the item is "
               "NOT eligible for installation on a type certificated aeroplane."
               if attrs.get("nonApprovedDesignData") else ""))
        heading = "Authorised Release Certificate"

    blocks = [
        Paragraph(heading, style["h"]),
        _kv_table(rows),
        Spacer(1, 6 * mm),
        Paragraph("Statement", style["h"]),
        Paragraph(statement, style["body"]),
    ]
    if record.get("notes"):
        blocks += [Spacer(1, 4 * mm),
                   Paragraph("Note on this record", style["h"]),
                   Paragraph(record["notes"], style["body"])]
    write_pdf(path, record["pn"], record["description"], blocks)


def build_bulletin_pdf(path: Path, record: dict):
    attrs = record["attrs"]
    style = _styles()
    blocks = [
        _kv_table([
            ("Service bulletin number", record["pn"]),
            ("Type", attrs.get("bulletinType", "-")),
            ("Category", attrs.get("bulletinCategory", "-")),
            ("Status", record["status"]),
            ("Publication date", attrs.get("publicationDate") or "not yet issued"),
            ("Effectivity", attrs.get("effectivityRange", "-")),
            ("Compliance period", attrs.get("compliancePeriod", "-")),
            ("Estimated labour", f"{attrs.get('estimatedManHours', '-')} hours"),
            ("Aircraft downtime", f"{attrs.get('aircraftDowntimeHours', '-')} hours"),
            ("Modification kit", attrs.get("kitPartNumber") or "none"),
            ("Weight change", f"{attrs.get('weightChangeKg', 0)} kg"),
            ("Related directive", attrs.get("relatedDirectiveRef") or "none"),
            ("Warranty", "applicable" if attrs.get("warrantyApplicable") else "not "
                                                                               "applicable"),
            ("Objective evidence", attrs.get("objectiveEvidenceRefs", "-")),
        ]),
        Spacer(1, 5 * mm),
        Paragraph("1. Planning information", style["h"]),
        Paragraph(f"<b>Effectivity.</b> {attrs.get('effectivityRange', '-')}",
                  style["body"]),
        Paragraph(f"<b>Compliance.</b> {attrs.get('compliancePeriod', '-')}",
                  style["body"]),
        Paragraph(
            "<b>Reason.</b> " + (record.get("notes") or
                                 "See the engineering change that originated this "
                                 "bulletin."), style["body"]),
        Spacer(1, 3 * mm),
        Paragraph("2. Material information", style["h"]),
        Paragraph(
            f"Modification kit {attrs.get('kitPartNumber') or 'not applicable'}. Parts "
            f"supplied under this bulletin are released on an authorised release "
            f"certificate; record the certificate number in the aeroplane records.",
            style["body"]),
        Spacer(1, 3 * mm),
        Paragraph("3. Accomplishment instructions", style["h"]),
        _grid_table(
            ["Step", "Action", "Reference"],
            [["1", "Make the aeroplane safe and gain access to the affected area.",
              "AMM chapter for the affected system"],
             ["2", "Remove and discard the superseded part. Record the part number and "
                   "serial or batch number removed.", "This bulletin, effectivity"],
             ["3", "Install the replacement part. Torque as specified.",
              "LSS-SPEC-PRC-004"],
             ["4", "Carry out the functional check for the affected system.",
              "AMM functional check"],
             ["5", "Record accomplishment, including the authorised release certificate "
                   "number of the part installed.", "Aeroplane records"],
             ["6", "Update the aeroplane configuration record.",
              "Operator configuration control"]],
            [14 * mm, 104 * mm, 52 * mm]),
        Spacer(1, 4 * mm),
        Paragraph("4. Why this bulletin exists", style["h"]),
        Paragraph(
            "A design change only reaches an aeroplane that has already been delivered "
            "through a service bulletin. The engineering change order makes the change "
            "real in production; this bulletin makes it real in the fleet. Where the "
            "authority issues an airworthiness directive against it, compliance stops "
            "being a recommendation and becomes a legal obligation on the operator.",
            style["body"]),
    ]
    write_pdf(path, record["pn"], record["description"], blocks)


def build_ncr_docx(path: Path, record: dict):
    attrs = record["attrs"]
    meta = [
        ("Nonconformance number", record["pn"]),
        ("Status", record["status"]),
        ("Severity", attrs.get("severity", "-")),
        ("Raised by", attrs.get("authorName", "-")),
        ("Raised on", attrs.get("issueDate", "-")),
        ("AS9100D clause", attrs.get("qmsClause", "-")),
        ("Owning function", attrs.get("recordOwnerFunction", "-")),
        ("Affected part", f"{attrs.get('affectedPartNumber', '-')} revision "
                          f"{attrs.get('affectedPartRevision', '-')}"),
        ("Affected serial / lot", attrs.get("affectedSerialNumbers", "-")),
        ("Quantity affected", attrs.get("quantityAffected", "-")),
        ("Detected at", attrs.get("detectedAt", "-")),
        ("Nonconformance type", attrs.get("nonconformanceType", "-")),
        ("Safety impact", "YES" if attrs.get("safetyImpact") else "No"),
        ("Escaped to customer", "YES" if attrs.get("escapeToCustomer") else "No"),
        ("Customer notified", attrs.get("customerNotificationDate") or "not applicable"),
        ("Prior recurrences", attrs.get("recurrenceCount", 0)),
        ("Cost of nonconformance",
         f"{attrs.get('costOfNonconformance', 0):,} (demo currency units)"),
    ]
    sections = [
        ("Description of the nonconformance", [record["description"]]),
        ("Containment", [attrs.get("containmentAction", "-"),
                         "Containment comes before disposition. The question 'where else "
                         "is this?' has to be answered before the question 'what do we "
                         "do with it?', because the second answer changes if the first "
                         "one is wide."]),
        ("Material review board disposition", [
            (["Item", "Entry"], [
                ["Disposition", attrs.get("mrbDisposition", "-")],
                ["Disposition authority", attrs.get("dispositionAuthority", "-")],
                ["Approved repair scheme", attrs.get("repairSchemeRef") or "none"],
            ]),
            "Safety impact and escape to the customer are recorded as two separate "
            "questions, because they are two separate questions. An escape is a "
            "customer notification matter; a safety impact is a continued airworthiness "
            "matter. One does not imply the other, and a record that conflates them "
            "will eventually get one of them wrong.",
        ]),
        ("Root cause and corrective action", [
            f"Root cause: {attrs.get('rootCause', '-')}",
            f"Corrective action reference: "
            f"{attrs.get('correctiveActionRef') or 'none raised'}",
            "The disposition deals with the parts. The corrective action deals with the "
            "cause. Closing the first without the second is how a nonconformance comes "
            "back, which is what the recurrence count above is for.",
        ]),
    ]
    if record.get("notes"):
        sections.append(("Note on this record", [record["notes"]]))
    write_docx(path, record["pn"], record["description"], meta, sections)


# ─────────────────────────────────────────────────────────────────────────────

def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 \
        else Path(__file__).parent.parent / "build" / "documents"
    out.mkdir(parents=True, exist_ok=True)

    written = 0
    skipped = []

    records = (DOC.ALL_DOCUMENTS + QA.FAI_REPORTS + QA.NCRS + QA.COCS + QA.ARCS
               + QA.SUPPLIER_RECORDS + QA.PROCESS_SPECS + QA.SERVICE_BULLETINS)

    for record in records:
        name = record.get("file")
        if not name:
            continue
        path = out / name
        stem = record["type"]
        try:
            if stem.endswith("Drawing"):
                build_drawing(path, record)
            elif stem.endswith("FaiReport"):
                build_fai_xlsx(path, record)
            elif stem.endswith("Ncr"):
                build_ncr_docx(path, record)
            elif stem.endswith("CofC"):
                build_certificate_pdf(path, record, "coc")
            elif stem.endswith("AirworthinessRelease"):
                build_certificate_pdf(path, record, "arc")
            elif stem.endswith("ServiceBulletin"):
                build_bulletin_pdf(path, record)
            elif stem.endswith("SupplierQualRecord"):
                build_supplier_xlsx(path, record)
            elif name.endswith(".xlsx"):
                build_checklist_xlsx(path, record)
            else:
                kind = stem.replace("n5LightSpeedSystems", "")
                build_generic_docx(path, record, kind, sections_for(record))
            written += 1
        except Exception as error:                 # noqa: BLE001
            skipped.append(f"{record['pn']} ({name}): {error}")

    print(f"Wrote {written} document files to {out}")
    by_ext: dict[str, int] = {}
    for path in sorted(out.iterdir()):
        by_ext[path.suffix] = by_ext.get(path.suffix, 0) + 1
    for ext in sorted(by_ext):
        print(f"  {by_ext[ext]:>3} {ext}")
    if skipped:
        print(f"\n{len(skipped)} not written:")
        for problem in skipped:
            print(f"  ! {problem}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
